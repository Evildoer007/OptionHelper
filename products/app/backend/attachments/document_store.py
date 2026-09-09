"""Immutable business-document storage and bounded text extraction."""

from __future__ import annotations

import base64
import binascii
import csv
import hashlib
import io
from itertools import islice
import json
import os
import re
import tempfile
from pathlib import Path
from typing import Any, Mapping, Sequence
import zipfile

from ..errors import ValidationError
from ..file_permissions import protect_private_path


_MEDIA_TYPES = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "text/csv": ".csv",
    "text/plain": ".txt",
    "text/markdown": ".md",
}
_ATTACHMENT_ID = re.compile(r"^sha256:[0-9a-f]{64}$")


class DocumentAttachmentStore:
    """Store source documents plus a bounded, searchable text projection."""

    max_documents_per_message = 20
    max_source_bytes = 20 * 1024 * 1024
    max_message_source_bytes = 200 * 1024 * 1024
    max_extracted_chars = 2_000_000
    max_sheet_rows = 20_000
    max_sheet_columns = 256
    max_pdf_pages = 1_000
    max_archive_entries = 10_000
    max_archive_uncompressed_bytes = 128 * 1024 * 1024
    max_archive_compression_ratio = 200

    def __init__(self, root: Path) -> None:
        self._root = Path(root).resolve() / "v1"
        self._mkdir_private(self._root)

    def save_encoded(self, documents: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
        if isinstance(documents, (str, bytes)) or not isinstance(documents, Sequence):
            raise ValidationError("文档附件必须是数组")
        if not documents or len(documents) > self.max_documents_per_message:
            raise ValidationError(f"每条消息最多添加{self.max_documents_per_message}个文档")
        decoded: list[tuple[bytes, str, str]] = []
        total = 0
        for item in documents:
            if not isinstance(item, Mapping) or set(item).difference({"data", "media_type", "name"}):
                raise ValidationError("文档附件字段无效")
            media_type = str(item.get("media_type", "")).strip().lower()
            if media_type not in _MEDIA_TYPES:
                raise ValidationError("仅支持PDF、DOCX、XLSX、CSV、TXT和Markdown文档")
            name = self._safe_name(item.get("name"))
            if not name or Path(name).suffix.lower() != _MEDIA_TYPES[media_type]:
                raise ValidationError("文档扩展名与声明格式不一致")
            encoded = item.get("data")
            if not isinstance(encoded, str) or not encoded:
                raise ValidationError("文档附件缺少数据")
            try:
                data = base64.b64decode(encoded, validate=True)
            except (binascii.Error, ValueError):
                raise ValidationError("文档附件不是有效的Base64数据") from None
            if not data or len(data) > self.max_source_bytes:
                raise ValidationError("单个文档超过20MB上限")
            total += len(data)
            if total > self.max_message_source_bytes:
                raise ValidationError("本条消息的文档总量超过200MB上限")
            decoded.append((data, media_type, name))
        return [self._commit(data, self._prepare(data, media_type, name)) for data, media_type, name in decoded]

    def read(self, reference: Mapping[str, Any]) -> bytes:
        digest = self._digest(str(reference.get("attachment_id", "")))
        source_path, text_path, metadata_path = self._paths(digest)
        for path in (source_path, text_path, metadata_path):
            self._reject_link(path)
        try:
            data = source_path.read_bytes()
            stored = json.loads(metadata_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            raise ValidationError("文档附件不可用") from None
        if hashlib.sha256(data).hexdigest() != digest or stored != dict(reference):
            raise ValidationError("文档附件完整性校验失败")
        return data

    def extracted_text(self, reference: Mapping[str, Any]) -> str:
        digest = self._digest(str(reference.get("attachment_id", "")))
        _source_path, text_path, _metadata_path = self._paths(digest)
        self.read(reference)
        self._reject_link(text_path)
        try:
            return text_path.read_text(encoding="utf-8")
        except (OSError, UnicodeError):
            raise ValidationError("文档文本索引不可用") from None

    def search(self, reference: Mapping[str, Any], query: str, *, limit: int = 8) -> list[dict[str, Any]]:
        terms = [term.casefold() for term in re.findall(r"[\w\u4e00-\u9fff]+", str(query)) if len(term) > 1]
        if not terms:
            raise ValidationError("文档检索词不能为空")
        chunks = self._chunks(self.extracted_text(reference))
        ranked: list[tuple[int, int, str]] = []
        for index, chunk in enumerate(chunks):
            folded = chunk.casefold()
            score = sum(folded.count(term) for term in terms)
            if score:
                ranked.append((score, index, chunk))
        ranked.sort(key=lambda item: (-item[0], item[1]))
        return [{"chunk": index, "text": text} for _score, index, text in ranked[:max(1, min(limit, 20))]]

    def read_chunk(self, reference: Mapping[str, Any], index: int) -> dict[str, Any]:
        chunks = self._chunks(self.extracted_text(reference))
        if index < 0 or index >= len(chunks):
            raise ValidationError("文档片段不存在")
        return {"chunk": index, "text": chunks[index], "total_chunks": len(chunks)}

    def _prepare(self, data: bytes, media_type: str, name: str) -> dict[str, Any]:
        text, status = self._extract(data, media_type)
        text = text.replace("\x00", "").strip()[: self.max_extracted_chars]
        digest = hashlib.sha256(data).hexdigest()
        return {
            "attachment_id": f"sha256:{digest}",
            "kind": "document",
            "media_type": media_type,
            "bytes": len(data),
            "name": name,
            "extraction_status": status,
            "text_chars": len(text),
            "_text": text,
        }

    def _extract(self, data: bytes, media_type: str) -> tuple[str, str]:
        if media_type in {"text/plain", "text/markdown", "text/csv"}:
            text = self._decode_text(data)
            if media_type == "text/csv":
                rows = csv.reader(io.StringIO(text))
                text = "\n".join(
                    "\t".join(cell.strip() for cell in row[: self.max_sheet_columns])
                    for row in islice(rows, self.max_sheet_rows)
                )
            return text, "ready"
        if media_type == "application/pdf":
            from pypdf import PdfReader
            try:
                reader = PdfReader(io.BytesIO(data))
                if len(reader.pages) > self.max_pdf_pages:
                    raise ValidationError("PDF页数超过1000页上限")
                parts: list[str] = []
                length = 0
                for page in reader.pages:
                    value = str(page.extract_text() or "")
                    parts.append(value)
                    length += len(value)
                    if length >= self.max_extracted_chars:
                        break
                text = "\n\n".join(parts)
            except ValidationError:
                raise
            except Exception as error:
                raise ValidationError("PDF文档无法解析") from error
            return (text, "ready") if text.strip() else ("", "no_text")
        if media_type.endswith("wordprocessingml.document"):
            from docx import Document
            try:
                self._validate_office_archive(data)
                document = Document(io.BytesIO(data))
                parts: list[str] = []
                length = 0
                for paragraph in document.paragraphs:
                    value = paragraph.text
                    parts.append(value)
                    length += len(value)
                    if length >= self.max_extracted_chars:
                        break
                for table in document.tables:
                    if length >= self.max_extracted_chars:
                        break
                    for row in table.rows:
                        value = "\t".join(cell.text for cell in row.cells)
                        parts.append(value)
                        length += len(value)
                        if length >= self.max_extracted_chars:
                            break
                return "\n".join(parts), "ready"
            except ValidationError:
                raise
            except Exception as error:
                raise ValidationError("DOCX文档无法解析") from error
        if media_type.endswith("spreadsheetml.sheet"):
            from openpyxl import load_workbook
            try:
                self._validate_office_archive(data)
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                parts: list[str] = []
                length = 0
                for sheet in workbook.worksheets:
                    heading = f"工作表：{sheet.title}"
                    parts.append(heading)
                    length += len(heading)
                    for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                        if row_index >= self.max_sheet_rows or length >= self.max_extracted_chars:
                            break
                        value = "\t".join("" if item is None else str(item) for item in row[: self.max_sheet_columns])
                        parts.append(value)
                        length += len(value)
                    if length >= self.max_extracted_chars:
                        break
                workbook.close()
                return "\n".join(parts), "ready"
            except ValidationError:
                raise
            except Exception as error:
                raise ValidationError("XLSX文档无法解析") from error
        raise ValidationError("文档格式未注册")

    def _validate_office_archive(self, data: bytes) -> None:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                if len(members) > self.max_archive_entries:
                    raise ValidationError("Office文档内部文件数量超过上限")
                total = 0
                for member in members:
                    if member.flag_bits & 0x1:
                        raise ValidationError("不支持加密Office文档")
                    total += int(member.file_size)
                    if total > self.max_archive_uncompressed_bytes:
                        raise ValidationError("Office文档解压后超过128MB上限")
                    compressed = max(1, int(member.compress_size))
                    if member.file_size > 1_048_576 and member.file_size / compressed > self.max_archive_compression_ratio:
                        raise ValidationError("Office文档压缩比例异常")
        except ValidationError:
            raise
        except (zipfile.BadZipFile, OSError):
            raise ValidationError("Office文档压缩结构无效") from None

    @staticmethod
    def _decode_text(data: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ValidationError("文本文件编码无法识别")

    def _commit(self, data: bytes, metadata: dict[str, Any]) -> dict[str, Any]:
        text = str(metadata.pop("_text"))
        digest = self._digest(str(metadata["attachment_id"]))
        source_path, text_path, metadata_path = self._paths(digest)
        self._mkdir_private(source_path.parent)
        for path in (source_path, text_path, metadata_path):
            self._reject_link(path)
        if source_path.exists() and text_path.exists() and metadata_path.exists():
            self.read(metadata)
            return dict(metadata)
        self._atomic_write(source_path, data)
        self._atomic_write(text_path, text.encode("utf-8"))
        self._atomic_write(metadata_path, json.dumps(metadata, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
        self.read(metadata)
        return dict(metadata)

    def _paths(self, digest: str) -> tuple[Path, Path, Path]:
        directory = self._root / digest[:2]
        return directory / f"{digest}.document", directory / f"{digest}.txt", directory / f"{digest}.json"

    @staticmethod
    def _chunks(text: str, size: int = 4_000, overlap: int = 400) -> list[str]:
        if not text:
            return []
        chunks: list[str] = []
        start = 0
        while start < len(text):
            chunks.append(text[start : start + size])
            if start + size >= len(text):
                break
            start += size - overlap
        return chunks

    @staticmethod
    def _safe_name(value: object) -> str | None:
        if not isinstance(value, str):
            return None
        name = Path(value.replace("\\", "/")).name.strip()
        name = "".join(char for char in name if char.isprintable())[:160]
        return name or None

    @staticmethod
    def _digest(attachment_id: str) -> str:
        if not _ATTACHMENT_ID.fullmatch(attachment_id):
            raise ValidationError("文档附件引用无效")
        return attachment_id.removeprefix("sha256:")

    @staticmethod
    def _mkdir_private(path: Path) -> None:
        path.mkdir(parents=True, exist_ok=True)
        if path.is_symlink():
            raise ValidationError("文档附件目录不能是符号链接")
        protect_private_path(path)

    @staticmethod
    def _reject_link(path: Path) -> None:
        if path.is_symlink():
            raise ValidationError("文档附件不能使用符号链接")

    @staticmethod
    def _atomic_write(path: Path, data: bytes) -> None:
        descriptor, temporary_name = tempfile.mkstemp(prefix=".pending-", dir=path.parent)
        temporary = Path(temporary_name)
        try:
            with os.fdopen(descriptor, "wb") as handle:
                protect_private_path(temporary)
                handle.write(data)
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temporary, path)
            protect_private_path(path)
        finally:
            if temporary.exists():
                temporary.unlink()
