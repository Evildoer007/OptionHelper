"""Render sanitised Designer HTML as a genuinely editable Word document.

The renderer consumes only the supplied HTML and frozen chart specifications.
It never reads ReportIR, local image paths, or network resources.  Paragraphs,
lists, tables and equations remain native Word content; charts are inserted as
static figures with native data tables so their current values remain visible.
"""

from __future__ import annotations

import base64
import binascii
from functools import lru_cache
from io import BytesIO
import math
from pathlib import Path
import re
from typing import Any, Mapping, Sequence
from urllib.parse import unquote_to_bytes
from xml.etree import ElementTree

from .design_tokens import TOKENS
from .system_fonts import windows_font_candidates

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image, ImageColor, ImageDraw, ImageFont


class WordRenderError(ValueError):
    """Raised when HTML cannot be represented safely as editable Word."""


_A4_WIDTH_MM = 210
_A4_HEIGHT_MM = 297
_CONTENT_WIDTH_INCHES = 6.65
_BODY_FONT_EAST_ASIA = "Songti SC"
_HEADING_FONT_EAST_ASIA = "Heiti SC"
_MAX_IMAGE_BYTES = 24 * 1024 * 1024
_MAX_IMAGE_PIXELS = 36_000_000
_RASTER_MIME_TYPES = {
    "image/png": "PNG",
    "image/jpeg": "JPEG",
    "image/jpg": "JPEG",
    "image/gif": "GIF",
    "image/webp": "WEBP",
    "image/bmp": "BMP",
    "image/tiff": "TIFF",
}
_BLOCK_TAGS = {
    "address", "article", "aside", "blockquote", "dd", "div", "dl", "dt",
    "figcaption", "figure", "footer", "h1", "h2", "h3", "h4", "h5", "h6",
    "header", "hr", "li", "main", "nav", "ol", "p", "pre", "section", "table", "ul",
}
_IGNORED_TAGS = {"head", "script", "style", "noscript", "template", "title"}
_CHART_PALETTE = ("#C8102E", "#315D8A", "#C28B2C", "#6B3FA0", "#2F7D6D")


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1].lower()


def _normalise_space(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\xa0", " "))


def _clean_text(value: str) -> str:
    return _normalise_space(value).strip()


def _positive_int(value: Any, fallback: int = 1) -> int:
    try:
        parsed = int(str(value))
    except (TypeError, ValueError):
        return fallback
    return parsed if 1 <= parsed <= 100 else fallback


def _parse_style(value: Any) -> dict[str, str]:
    return {
        name.strip().lower(): setting.strip()
        for name, setting in re.findall(r"([\w-]+)\s*:\s*([^;]+)", str(value or ""))
    }


def _parse_points(value: str | None, fallback: float | None = None) -> float | None:
    candidate = str(value or "").strip().lower()
    match = re.fullmatch(r"([-+]?(?:\d+(?:\.\d*)?|\.\d+))(pt|px|em|rem|%)?", candidate)
    if not match:
        return fallback
    number = float(match.group(1))
    unit = match.group(2) or "pt"
    if unit == "px":
        number *= 0.75
    elif unit in {"em", "rem"}:
        number *= fallback or 11.0
    elif unit == "%":
        number = (fallback or 11.0) * number / 100
    return min(72.0, max(5.0, number))


def _parse_color(value: Any) -> RGBColor | None:
    candidate = str(value or "").strip()
    short = re.fullmatch(r"#([0-9A-Fa-f]{3})", candidate)
    if short:
        candidate = "#" + "".join(char * 2 for char in short.group(1))
    full = re.fullmatch(r"#([0-9A-Fa-f]{6})", candidate)
    if full:
        return RGBColor.from_string(full.group(1).upper())
    rgb = re.fullmatch(r"rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)(?:\s*,[^)]*)?\)", candidate)
    if rgb:
        channels = tuple(min(255, max(0, int(item))) for item in rgb.groups())
        return RGBColor(*channels)
    return None


def _set_font(run: Any, *, east_asia: str = _BODY_FONT_EAST_ASIA, latin: str = "Arial") -> None:
    run.font.name = latin
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)


def _set_style_font(style: Any, *, east_asia: str, latin: str, size: float, bold: bool = False) -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    properties = style.element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for name in ("asciiTheme", "eastAsiaTheme", "hAnsiTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{name}"), None)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)


def _configure_document(document: DocumentType) -> None:
    section = document.sections[0]
    section.page_width = Mm(_A4_WIDTH_MM)
    section.page_height = Mm(_A4_HEIGHT_MM)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)

    styles = document.styles
    _set_style_font(styles["Normal"], east_asia=_BODY_FONT_EAST_ASIA, latin="Arial", size=10.5)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.35
    for name, east_asia, size, bold in (
        ("Title", _HEADING_FONT_EAST_ASIA, 20, True),
        ("Heading 1", _HEADING_FONT_EAST_ASIA, 15, True),
        ("Heading 2", _HEADING_FONT_EAST_ASIA, 13, True),
        ("Heading 3", _HEADING_FONT_EAST_ASIA, 11.5, True),
        ("Caption", _BODY_FONT_EAST_ASIA, 9, False),
    ):
        _set_style_font(styles[name], east_asia=east_asia, latin="Arial", size=size, bold=bold)
    styles["Title"].paragraph_format.space_after = Pt(12)
    title_properties = styles["Title"].element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    styles["Heading 3"].paragraph_format.space_before = Pt(7)
    styles["Heading 3"].paragraph_format.space_after = Pt(3)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        styles[name].paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("第 ")
    _set_font(footer_run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = OxmlElement("w:r")
    field_run.extend((begin, instruction, separate, value, end))
    footer._p.append(field_run)
    tail = footer.add_run(" 页")
    _set_font(tail)

    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    document.settings._element.append(update_fields)


def _paragraph_alignment(node: Tag) -> WD_ALIGN_PARAGRAPH | None:
    style = _parse_style(node.get("style"))
    value = str(node.get("align") or style.get("text-align") or "").strip().lower()
    classes = set(node.get("class") or ())
    if not value:
        if {"text-center", "center", "report-title--center"} & classes:
            value = "center"
        elif {"text-right", "right"} & classes:
            value = "right"
        elif {"text-justify", "justify"} & classes:
            value = "justify"
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(value)


def _context_for(node: Tag, inherited: Mapping[str, Any]) -> dict[str, Any]:
    context = dict(inherited)
    tag = node.name.lower()
    style = _parse_style(node.get("style"))
    if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        context.setdefault("east_asia", _HEADING_FONT_EAST_ASIA)
    if tag in {"strong", "b"} or style.get("font-weight", "").lower() in {"bold", "600", "700", "800", "900"}:
        context["bold"] = True
    if tag in {"em", "i"} or style.get("font-style", "").lower() == "italic":
        context["italic"] = True
    decoration = style.get("text-decoration", "").lower()
    if tag == "u" or "underline" in decoration:
        context["underline"] = True
    if tag in {"s", "strike", "del"} or "line-through" in decoration:
        context["strike"] = True
    if tag == "sup":
        context["superscript"] = True
    if tag == "sub":
        context["subscript"] = True
    if {"conclusion-band__label", "conclusion-heading"} & set(node.get("class") or ()):
        context["color"] = _parse_color(TOKENS.colors["gold_ink"])
    color = _parse_color(style.get("color"))
    if color is not None:
        context["color"] = color
    size = _parse_points(style.get("font-size"), context.get("size", 10.5))
    if size is not None and "font-size" in style:
        context["size"] = size
    family = style.get("font-family")
    if family:
        names = [item.strip().strip("\"'") for item in family.split(",")]
        if names:
            context["east_asia"] = names[0]
            context["latin"] = next((item for item in names if item.isascii()), context.get("latin", "Arial"))
    return context


def _apply_run_context(run: Any, context: Mapping[str, Any]) -> None:
    _set_font(run, east_asia=str(context.get("east_asia") or _BODY_FONT_EAST_ASIA), latin=str(context.get("latin") or "Arial"))
    run.bold = bool(context.get("bold"))
    run.italic = bool(context.get("italic"))
    run.underline = bool(context.get("underline"))
    run.font.strike = bool(context.get("strike"))
    run.font.superscript = bool(context.get("superscript"))
    run.font.subscript = bool(context.get("subscript"))
    if context.get("color") is not None:
        run.font.color.rgb = context["color"]
    if context.get("size") is not None:
        run.font.size = Pt(float(context["size"]))


def _append_math_text(parent: Any, value: str) -> None:
    text = _normalise_space(value)
    if not text:
        return
    run = OxmlElement("m:r")
    node = OxmlElement("m:t")
    if text.startswith(" ") or text.endswith(" "):
        node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    parent.append(run)


def _math_children(node: Tag) -> list[Any]:
    return [child for child in node.children if not isinstance(child, NavigableString) or child.strip()]


def _append_math_nodes(parent: Any, nodes: Sequence[Any]) -> None:
    for node in nodes:
        if isinstance(node, NavigableString):
            _append_math_text(parent, str(node))
            continue
        if not isinstance(node, Tag):
            continue
        tag = node.name.lower()
        children = _math_children(node)
        if tag in {"math", "mrow", "mstyle", "mpadded", "mphantom"}:
            _append_math_nodes(parent, children)
        elif tag in {"mi", "mn", "mo", "mtext", "ms"}:
            _append_math_text(parent, node.get_text("", strip=False))
        elif tag == "mfrac" and len(children) >= 2:
            fraction = OxmlElement("m:f")
            numerator = OxmlElement("m:num")
            denominator = OxmlElement("m:den")
            _append_math_nodes(numerator, [children[0]])
            _append_math_nodes(denominator, [children[1]])
            fraction.extend((numerator, denominator))
            parent.append(fraction)
        elif tag in {"msub", "msup"} and len(children) >= 2:
            script = OxmlElement("m:sSub" if tag == "msub" else "m:sSup")
            base = OxmlElement("m:e")
            attachment = OxmlElement("m:sub" if tag == "msub" else "m:sup")
            _append_math_nodes(base, [children[0]])
            _append_math_nodes(attachment, [children[1]])
            script.extend((base, attachment))
            parent.append(script)
        elif tag == "msubsup" and len(children) >= 3:
            script = OxmlElement("m:sSubSup")
            base = OxmlElement("m:e")
            subscript = OxmlElement("m:sub")
            superscript = OxmlElement("m:sup")
            _append_math_nodes(base, [children[0]])
            _append_math_nodes(subscript, [children[1]])
            _append_math_nodes(superscript, [children[2]])
            script.extend((base, subscript, superscript))
            parent.append(script)
        elif tag in {"msqrt", "mroot"} and children:
            radical = OxmlElement("m:rad")
            properties = OxmlElement("m:radPr")
            degree = OxmlElement("m:deg")
            radicand = OxmlElement("m:e")
            if tag == "msqrt":
                hidden = OxmlElement("m:degHide")
                hidden.set(qn("m:val"), "1")
                properties.append(hidden)
                _append_math_nodes(radicand, children)
            else:
                if len(children) >= 2:
                    _append_math_nodes(degree, [children[1]])
                _append_math_nodes(radicand, [children[0]])
            radical.extend((properties, degree, radicand))
            parent.append(radical)
        elif tag == "mfenced":
            _append_math_text(parent, str(node.get("open", "(")))
            separator = str(node.get("separators", ","))[:1] or ","
            for index, child in enumerate(children):
                if index:
                    _append_math_text(parent, separator)
                _append_math_nodes(parent, [child])
            _append_math_text(parent, str(node.get("close", ")")))
        elif tag in {"mover", "munder"} and len(children) >= 2:
            limit = OxmlElement("m:limUpp" if tag == "mover" else "m:limLow")
            base = OxmlElement("m:e")
            attachment = OxmlElement("m:lim")
            _append_math_nodes(base, [children[0]])
            _append_math_nodes(attachment, [children[1]])
            limit.extend((base, attachment))
            parent.append(limit)
        elif tag == "munderover" and len(children) >= 3:
            upper = OxmlElement("m:limUpp")
            lower = OxmlElement("m:limLow")
            base = OxmlElement("m:e")
            lower_limit = OxmlElement("m:lim")
            upper_limit = OxmlElement("m:lim")
            _append_math_nodes(base, [children[0]])
            _append_math_nodes(lower_limit, [children[1]])
            lower.extend((base, lower_limit))
            upper_base = OxmlElement("m:e")
            upper_base.append(lower)
            _append_math_nodes(upper_limit, [children[2]])
            upper.extend((upper_base, upper_limit))
            parent.append(upper)
        elif tag == "mtable":
            matrix = OxmlElement("m:m")
            for row_node in (child for child in children if isinstance(child, Tag) and child.name.lower() == "mtr"):
                row = OxmlElement("m:mr")
                for cell_node in (child for child in _math_children(row_node) if isinstance(child, Tag) and child.name.lower() == "mtd"):
                    cell = OxmlElement("m:e")
                    _append_math_nodes(cell, _math_children(cell_node))
                    row.append(cell)
                matrix.append(row)
            parent.append(matrix)
        elif tag == "semantics":
            semantic = next((child for child in children if not isinstance(child, Tag) or child.name.lower() not in {"annotation", "annotation-xml"}), None)
            if semantic is not None:
                _append_math_nodes(parent, [semantic])
        elif tag not in {"annotation", "annotation-xml"}:
            _append_math_nodes(parent, children)


def _math_element(node: Tag, *, display: bool) -> Any:
    equation = OxmlElement("m:oMath")
    _append_math_nodes(equation, _math_children(node))
    if not display:
        return equation
    container = OxmlElement("m:oMathPara")
    container.append(equation)
    return container


def _set_repeat_table_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _set_cell_shading(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color.removeprefix("#").upper())


def _set_table_borders(table: Any, color: str = "D9D9D9") -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)
        borders.append(border)


def _set_cell_margins(cell: Any, *, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = margins.find(qn(f"w:{name}"))
        if item is None:
            item = OxmlElement(f"w:{name}")
            margins.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def _clear_paragraph(paragraph: Any) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    return parsed if math.isfinite(parsed) else None


def _compact_number(value: float) -> str:
    if value == 0:
        return "0"
    if abs(value) >= 1_000_000 or abs(value) < 0.0001:
        return f"{value:.4g}"
    return f"{value:.6f}".rstrip("0").rstrip(".")


def _display_value(value: Any, value_format: Any = None, suffix: Any = None) -> str:
    number = _number(value)
    if number is None:
        text = "—" if value is None else str(value)
    elif str(value_format or "").lower() == "percent":
        text = _compact_number(number * 100) + "%"
    else:
        text = _compact_number(number)
    return text + str(suffix or "")


def _font_candidates() -> tuple[Path, ...]:
    return (
        Path("/System/Library/Fonts/STHeiti Medium.ttc"),
        Path("/System/Library/Fonts/STHeiti Light.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
        *windows_font_candidates("msyh.ttc", "simsun.ttc"),
    )


@lru_cache(maxsize=16)
def _image_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for candidate in _font_candidates():
        if candidate.is_file():
            try:
                return ImageFont.truetype(str(candidate), size=size, index=0)
            except OSError:
                continue
    return ImageFont.load_default()


def _interpolate_color(start: str, end: str, ratio: float) -> str:
    ratio = max(0.0, min(1.0, ratio))
    left = ImageColor.getrgb(start)
    right = ImageColor.getrgb(end)
    channels = tuple(round(a + (b - a) * ratio) for a, b in zip(left, right))
    return "#" + "".join(f"{channel:02X}" for channel in channels)


def _render_chart_png(spec: Mapping[str, Any]) -> bytes:
    chart_type = str(spec.get("type") or "line").lower()
    if chart_type not in {"line", "bar", "heatmap"}:
        raise WordRenderError(f"Word不支持图表类型：{chart_type}。")
    width, height = 1280, 420
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    label_font = _image_font(22)
    small_font = _image_font(18)
    left, right, top, bottom = 110, 48, 56, 78
    plot_width = width - left - right
    plot_height = height - top - bottom
    draw.rectangle((left, top, left + plot_width, top + plot_height), outline="#B9BEC7", width=2)
    x_values = list(spec.get("x") or [])
    value_format = spec.get("value_format")
    suffix = spec.get("value_suffix")
    maximum_axis_labels = 8

    def draw_text_within_canvas(
        position: tuple[float, float],
        text: str,
        *,
        fill: str,
        font: ImageFont.FreeTypeFont | ImageFont.ImageFont,
    ) -> None:
        box = draw.textbbox((0, 0), text, font=font)
        minimum_x, maximum_x = 2 - box[0], width - 2 - box[2]
        minimum_y, maximum_y = 2 - box[1], height - 2 - box[3]
        x = min(max(position[0], minimum_x), maximum_x) if minimum_x <= maximum_x else minimum_x
        y = min(max(position[1], minimum_y), maximum_y) if minimum_y <= maximum_y else minimum_y
        draw.text((x, y), text, fill=fill, font=font)

    def sampled_indices(length: int) -> list[int]:
        if length <= maximum_axis_labels:
            return list(range(length))
        return sorted({round(position * (length - 1) / (maximum_axis_labels - 1)) for position in range(maximum_axis_labels)})

    if chart_type == "heatmap":
        y_values = list(spec.get("y") or [])
        data = [row for row in list(spec.get("data") or []) if isinstance(row, (list, tuple)) and len(row) == 3 and _number(row[2]) is not None]
        values = [float(row[2]) for row in data]
        low, high = (min(values), max(values)) if values else (0.0, 1.0)
        columns = max(1, len(x_values))
        rows = max(1, len(y_values))
        cell_width = plot_width / columns
        cell_height = plot_height / rows
        cells = {(int(row[0]), int(row[1])): float(row[2]) for row in data if str(row[0]).lstrip("-").isdigit() and str(row[1]).lstrip("-").isdigit()}
        for y_index in range(rows):
            for x_index in range(columns):
                value = cells.get((x_index, y_index))
                ratio = 0.5 if high == low else (value - low) / (high - low) if value is not None else 0.0
                color = "#FFFFFF" if value is None else _interpolate_color("#E8EEF5", "#C8102E", ratio)
                x0 = left + x_index * cell_width
                y0 = top + (rows - y_index - 1) * cell_height
                outline = "#D9DDE3" if value is None else "#FFFFFF"
                draw.rectangle((x0, y0, x0 + cell_width, y0 + cell_height), fill=color, outline=outline, width=2)
                if value is not None:
                    text = _display_value(value, value_format, suffix)
                    box = draw.textbbox((0, 0), text, font=small_font)
                    draw_text_within_canvas(
                        (x0 + (cell_width - (box[2] - box[0])) / 2, y0 + (cell_height - (box[3] - box[1])) / 2),
                        text,
                        fill="#FFFFFF" if ratio > 0.55 else "#252B35",
                        font=small_font,
                    )
        for index in sampled_indices(len(x_values)):
            value = x_values[index]
            text = str(value)
            box = draw.textbbox((0, 0), text, font=small_font)
            draw_text_within_canvas(
                (left + (index + 0.5) * cell_width - (box[2] - box[0]) / 2, top + plot_height + 14),
                text,
                fill="#4F5967",
                font=small_font,
            )
        for index in sampled_indices(len(y_values)):
            value = y_values[index]
            text = str(value)
            box = draw.textbbox((0, 0), text, font=small_font)
            draw_text_within_canvas(
                (left - (box[2] - box[0]) - 12, top + (rows - index - 0.5) * cell_height - (box[3] - box[1]) / 2),
                text,
                fill="#4F5967",
                font=small_font,
            )
    else:
        series = [item for item in list(spec.get("series") or []) if isinstance(item, Mapping)]
        numeric = [_number(value) for item in series for value in list(item.get("data") or [])]
        values = [value for value in numeric if value is not None]
        low, high = (min(values), max(values)) if values else (0.0, 1.0)
        if chart_type == "bar":
            low, high = min(0.0, low), max(0.0, high)
        if low == high:
            low, high = low - 1.0, high + 1.0
        for tick in range(5):
            y = top + plot_height - plot_height * tick / 4
            draw.line((left, y, left + plot_width, y), fill="#E1E4E8", width=1)
            value = low + (high - low) * tick / 4
            label = _display_value(value, value_format, suffix)
            box = draw.textbbox((0, 0), label, font=small_font)
            draw_text_within_canvas(
                (left - (box[2] - box[0]) - 12, y - (box[3] - box[1]) / 2),
                label,
                fill="#6F7780",
                font=small_font,
            )
        count = max(1, len(x_values))
        zero_y = top + plot_height - (0.0 - low) / (high - low) * plot_height
        for series_index, item in enumerate(series):
            color = _CHART_PALETTE[series_index % len(_CHART_PALETTE)]
            item_values = list(item.get("data") or [])
            points: list[tuple[int, float, float, float]] = []
            for index, raw in enumerate(item_values[:count]):
                value = _number(raw)
                if value is None:
                    continue
                x = left + (index + 0.5) * plot_width / count
                y = top + plot_height - (value - low) / (high - low) * plot_height
                points.append((index, x, y, value))
            if chart_type == "bar":
                group_width = plot_width / count
                bar_width = group_width * 0.72 / max(1, len(series))
                for _index, x, y, value in points:
                    shifted = x + (series_index - (len(series) - 1) / 2) * bar_width
                    draw.rectangle((shifted - bar_width * 0.42, min(zero_y, y), shifted + bar_width * 0.42, max(zero_y, y)), fill=color)
                    label = _display_value(value, value_format, suffix)
                    box = draw.textbbox((0, 0), label, font=small_font)
                    draw_text_within_canvas(
                        (shifted - (box[2] - box[0]) / 2, min(zero_y, y) - 24),
                        label,
                        fill=color,
                        font=small_font,
                    )
            elif points:
                segments: list[list[tuple[int, float, float, float]]] = []
                for point in points:
                    if not segments or point[0] != segments[-1][-1][0] + 1:
                        segments.append([point])
                    else:
                        segments[-1].append(point)
                for segment in segments:
                    if len(segment) > 1:
                        draw.line([(x, y) for _index, x, y, _value in segment], fill=color, width=4, joint="curve")
                show_markers = count <= 60
                show_value_labels = count <= maximum_axis_labels
                for _index, x, y, value in points:
                    if show_markers:
                        draw.ellipse((x - 6, y - 6, x + 6, y + 6), fill=color)
                    if show_value_labels:
                        draw_text_within_canvas((x + 8, y - 28), _display_value(value, value_format, suffix), fill=color, font=small_font)
            legend_x = left + 200 + series_index * 210
            draw.rectangle((legend_x, 29, legend_x + 22, 37), fill=color)
            draw_text_within_canvas(
                (legend_x + 30, 16),
                str(item.get("name") or f"序列{series_index + 1}"),
                fill="#4F5967",
                font=small_font,
            )
        for index in sampled_indices(len(x_values)):
            value = x_values[index]
            text = str(value)
            box = draw.textbbox((0, 0), text, font=small_font)
            x = left + (index + 0.5) * plot_width / count
            draw_text_within_canvas(
                (x - (box[2] - box[0]) / 2, top + plot_height + 14),
                text,
                fill="#4F5967",
                font=small_font,
            )

    x_axis_name = _clean_text(str(spec.get("x_axis_name") or ""))
    y_axis_name = _clean_text(str(spec.get("y_axis_name") or ""))
    if x_axis_name:
        box = draw.textbbox((0, 0), x_axis_name, font=label_font)
        draw_text_within_canvas(
            (left + (plot_width - (box[2] - box[0])) / 2, height - 38),
            x_axis_name,
            fill="#4F5967",
            font=label_font,
        )
    if y_axis_name:
        draw_text_within_canvas((left, 33), y_axis_name, fill="#4F5967", font=small_font)
    output = BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


def _data_uri_bytes(source: str) -> tuple[str, bytes]:
    if not source.lower().startswith("data:"):
        raise WordRenderError("Word图片仅接受data URI，不允许网络地址或本地文件路径。")
    header, separator, payload = source.partition(",")
    if not separator:
        raise WordRenderError("图片data URI格式无效。")
    metadata = header[5:].split(";")
    mime_type = metadata[0].strip().lower()
    try:
        data = base64.b64decode(payload, validate=True) if any(item.lower() == "base64" for item in metadata[1:]) else unquote_to_bytes(payload)
    except (binascii.Error, ValueError) as error:
        raise WordRenderError("图片data URI编码无效。") from error
    if len(data) > _MAX_IMAGE_BYTES:
        raise WordRenderError("图片data URI超过允许大小。")
    return mime_type, data


def _raster_to_png(data: bytes, expected_format: str) -> tuple[bytes, tuple[int, int]]:
    try:
        with Image.open(BytesIO(data)) as source:
            if source.width * source.height > _MAX_IMAGE_PIXELS:
                raise WordRenderError("图片像素尺寸超过允许范围。")
            source.load()
            if str(source.format or "").upper() != expected_format:
                raise WordRenderError("图片data URI声明格式与实际内容不一致。")
            image = source.convert("RGBA")
    except WordRenderError:
        raise
    except Exception as error:
        raise WordRenderError("不支持的图片格式或图片data URI内容损坏，无法读取。") from error
    output = BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue(), image.size


def _svg_number(value: Any, fallback: float = 0.0) -> float:
    match = re.search(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)", str(value or ""))
    return float(match.group(0)) if match else fallback


def _svg_styles(root: ElementTree.Element) -> dict[str, dict[str, str]]:
    styles: dict[str, dict[str, str]] = {}
    for element in root.iter():
        if _local_name(element.tag) != "style":
            continue
        for class_name, declaration in re.findall(r"\.([\w-]+)\s*\{([^}]*)\}", element.text or ""):
            styles[class_name] = _parse_style(declaration)
    return styles


def _svg_paint(element: ElementTree.Element, styles: Mapping[str, Mapping[str, str]], inherited: Mapping[str, str]) -> dict[str, str]:
    paint = dict(inherited)
    for class_name in (element.get("class") or "").split():
        paint.update(styles.get(class_name, {}))
    for name in ("fill", "stroke", "stroke-width", "font-size", "font-weight", "opacity", "fill-opacity", "stroke-opacity", "text-anchor"):
        if element.get(name) is not None:
            paint[name] = str(element.get(name))
    paint.update(_parse_style(element.get("style")))
    return paint


def _svg_color(value: Any, fallback: str | None = None) -> tuple[int, int, int, int] | None:
    candidate = str(value or fallback or "").strip()
    if not candidate or candidate.lower() in {"none", "transparent"}:
        return None
    if candidate.startswith("url("):
        candidate = "#C8102E"
    try:
        rgb = ImageColor.getrgb(candidate)
    except ValueError:
        if fallback is None:
            return None
        rgb = ImageColor.getrgb(fallback)
    return (*rgb[:3], 255)


def _svg_path_commands(raw: str) -> list[tuple[str, list[float]]]:
    tokens = re.findall(r"[MmLlHhVvCcQqZz]|[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?", raw)
    sizes = {"M": 2, "L": 2, "H": 1, "V": 1, "C": 6, "Q": 4, "Z": 0}
    result: list[tuple[str, list[float]]] = []
    index = 0
    active = ""
    while index < len(tokens):
        if tokens[index].isalpha():
            active = tokens[index]
            index += 1
            if active.upper() == "Z":
                result.append((active, []))
                active = ""
                continue
        count = sizes.get(active.upper()) if active else None
        if count is None or index + count > len(tokens) or any(item.isalpha() for item in tokens[index:index + count]):
            break
        result.append((active, [float(item) for item in tokens[index:index + count]]))
        index += count
        if active in {"M", "m"}:
            active = "L" if active == "M" else "l"
    return result


def _svg_to_png_with_svglib(data: bytes) -> tuple[bytes, tuple[int, int]] | None:
    """Use the packaged svglib/ReportLab path when its raster backend exists."""

    try:
        from svglib.svglib import svg2rlg
        from reportlab.graphics import renderPM
    except ImportError:
        return None
    try:
        drawing = svg2rlg(BytesIO(data))
        if drawing is None or not drawing.width or not drawing.height:
            return None
        scale = min(1.5, 1400 / float(drawing.width), 900 / float(drawing.height))
        drawing.scale(scale, scale)
        drawing.width *= scale
        drawing.height *= scale
        png = renderPM.drawToString(drawing, fmt="PNG")
        return png, (max(1, round(drawing.width)), max(1, round(drawing.height)))
    except Exception:
        return None


def _svg_to_png(data: bytes) -> tuple[bytes, tuple[int, int]]:
    # Validate before delegating: optional SVG engines may fetch linked images,
    # stylesheets, or local files even when the outer image is a data URI.
    try:
        root = ElementTree.fromstring(data)
    except ElementTree.ParseError as error:
        raise WordRenderError("SVG图片格式无效。") from error
    if _local_name(root.tag) != "svg":
        raise WordRenderError("SVG图片缺少svg根元素。")
    for element in root.iter():
        for key, value in element.attrib.items():
            name = _local_name(key)
            if name in {"href", "src", "base"} and (name == "base" or not re.fullmatch(r"#[A-Za-z_][A-Za-z0-9_.:-]*", value.strip())):
                raise WordRenderError("SVG图片不允许引用外部或本地资源。")
        styles = [str(value) for value in element.attrib.values()]
        if _local_name(element.tag) == "style":
            styles.append("".join(element.itertext()))
        for style in styles:
            if "\\" in style or "@" in style or re.search(r"(?:image-set|src)\s*\(", style, re.IGNORECASE):
                raise WordRenderError("SVG图片包含不受支持的资源样式。")
            for match in re.finditer(r"url\s*\((.*?)\)", style, re.IGNORECASE | re.DOTALL):
                if not re.fullmatch(r"#[A-Za-z_][A-Za-z0-9_.:-]*", match.group(1).strip().strip("\"'")):
                    raise WordRenderError("SVG图片不允许引用外部或本地资源。")
    converted = _svg_to_png_with_svglib(data)
    if converted is not None:
        return converted
    view_box = str(root.get("viewBox") or root.get("viewbox") or "").replace(",", " ").split()
    if len(view_box) == 4:
        offset_x, offset_y, source_width, source_height = (float(value) for value in view_box)
    else:
        offset_x = offset_y = 0.0
        source_width = _svg_number(root.get("width"), 1200.0)
        source_height = _svg_number(root.get("height"), 600.0)
    if source_width <= 0 or source_height <= 0:
        raise WordRenderError("SVG图片缺少有效画布尺寸。")
    scale = min(1.5, 1400 / source_width, 900 / source_height)
    width = max(1, round(source_width * scale))
    height = max(1, round(source_height * scale))
    image = Image.new("RGBA", (width, height), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    styles = _svg_styles(root)

    def point(x: float, y: float) -> tuple[float, float]:
        return ((x - offset_x) * scale, (y - offset_y) * scale)

    def paint_color(paint: Mapping[str, str], name: str, fallback: str | None = None) -> tuple[int, int, int, int] | None:
        color = _svg_color(paint.get(name), fallback)
        if color is None:
            return None
        opacity = _svg_number(paint.get("opacity"), 1.0) * _svg_number(paint.get(f"{name}-opacity"), 1.0)
        return (*color[:3], round(255 * max(0.0, min(1.0, opacity))))

    def draw_path(raw: str, paint: Mapping[str, str]) -> None:
        subpaths: list[tuple[list[tuple[float, float]], bool]] = []
        points: list[tuple[float, float]] = []
        current = (0.0, 0.0)
        origin = (0.0, 0.0)
        closed = False

        def flush() -> None:
            nonlocal points, closed
            if points:
                subpaths.append((points, closed))
            points = []
            closed = False

        for command, values in _svg_path_commands(raw):
            absolute = command.isupper()
            name = command.upper()
            if name == "M":
                flush()
                x, y = values
                current = (x, y) if absolute else (current[0] + x, current[1] + y)
                origin = current
                points.append(point(*current))
            elif name == "L":
                x, y = values
                current = (x, y) if absolute else (current[0] + x, current[1] + y)
                points.append(point(*current))
            elif name == "H":
                current = (values[0] if absolute else current[0] + values[0], current[1])
                points.append(point(*current))
            elif name == "V":
                current = (current[0], values[0] if absolute else current[1] + values[0])
                points.append(point(*current))
            elif name in {"C", "Q"}:
                start = current
                coordinates = [(values[index], values[index + 1]) for index in range(0, len(values), 2)]
                if not absolute:
                    coordinates = [(start[0] + x, start[1] + y) for x, y in coordinates]
                end = coordinates[-1]
                steps = 16
                for step in range(1, steps + 1):
                    t = step / steps
                    if name == "C":
                        c1, c2 = coordinates[0], coordinates[1]
                        x = (1 - t) ** 3 * start[0] + 3 * (1 - t) ** 2 * t * c1[0] + 3 * (1 - t) * t ** 2 * c2[0] + t ** 3 * end[0]
                        y = (1 - t) ** 3 * start[1] + 3 * (1 - t) ** 2 * t * c1[1] + 3 * (1 - t) * t ** 2 * c2[1] + t ** 3 * end[1]
                    else:
                        control = coordinates[0]
                        x = (1 - t) ** 2 * start[0] + 2 * (1 - t) * t * control[0] + t ** 2 * end[0]
                        y = (1 - t) ** 2 * start[1] + 2 * (1 - t) * t * control[1] + t ** 2 * end[1]
                    points.append(point(x, y))
                current = end
            elif name == "Z":
                closed = True
                current = origin
        flush()
        fill = paint_color(paint, "fill")
        stroke = paint_color(paint, "stroke")
        stroke_width = max(1, round(_svg_number(paint.get("stroke-width"), 1.0) * scale))
        for path_points, is_closed in subpaths:
            if fill is not None and is_closed and len(path_points) >= 3:
                draw.polygon(path_points, fill=fill)
            if stroke is not None and len(path_points) >= 2:
                draw.line(path_points + ([path_points[0]] if is_closed else []), fill=stroke, width=stroke_width, joint="curve")

    def walk(element: ElementTree.Element, inherited: Mapping[str, str]) -> None:
        tag = _local_name(element.tag)
        if tag in {"defs", "style", "title", "desc", "lineargradient", "radialgradient", "stop", "clippath"}:
            return
        if element.get("transform"):
            raise WordRenderError("SVG图片包含暂不支持的transform变换。")
        if tag in {"image", "use", "foreignobject", "filter", "mask", "pattern"}:
            raise WordRenderError(f"SVG图片包含不支持的{tag}元素。")
        paint = _svg_paint(element, styles, inherited)
        fill = paint_color(paint, "fill", "#000000")
        stroke = paint_color(paint, "stroke")
        stroke_width = max(1, round(_svg_number(paint.get("stroke-width"), 1.0) * scale))
        if tag == "rect":
            x, y = point(_svg_number(element.get("x")), _svg_number(element.get("y")))
            w = _svg_number(element.get("width")) * scale
            h = _svg_number(element.get("height")) * scale
            draw.rectangle((x, y, x + w, y + h), fill=fill, outline=stroke, width=stroke_width)
        elif tag == "line" and stroke is not None:
            draw.line((*point(_svg_number(element.get("x1")), _svg_number(element.get("y1"))), *point(_svg_number(element.get("x2")), _svg_number(element.get("y2")))), fill=stroke, width=stroke_width)
        elif tag in {"circle", "ellipse"}:
            cx, cy = point(_svg_number(element.get("cx")), _svg_number(element.get("cy")))
            rx = _svg_number(element.get("r") or element.get("rx")) * scale
            ry = _svg_number(element.get("r") or element.get("ry")) * scale
            draw.ellipse((cx - rx, cy - ry, cx + rx, cy + ry), fill=fill, outline=stroke, width=stroke_width)
        elif tag in {"path", "polyline", "polygon"}:
            raw = element.get("d") or ""
            if tag != "path":
                numbers = [float(value) for value in re.findall(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)", element.get("points") or "")]
                raw = "M" + " L".join(f"{numbers[index]},{numbers[index + 1]}" for index in range(0, len(numbers) - 1, 2))
                if tag == "polygon":
                    raw += " Z"
            if raw:
                draw_path(raw, paint)
        elif tag == "text":
            content = _clean_text("".join(element.itertext()))
            if content:
                font_size = max(7, round(_svg_number(paint.get("font-size"), 12.0) * scale))
                font = _image_font(font_size)
                x, y = point(_svg_number(element.get("x")), _svg_number(element.get("y")))
                anchor = str(paint.get("text-anchor") or element.get("text-anchor") or "start")
                box = draw.textbbox((0, 0), content, font=font)
                if anchor == "middle":
                    x -= (box[2] - box[0]) / 2
                elif anchor == "end":
                    x -= box[2] - box[0]
                draw.text((x, y - (box[3] - box[1])), content, fill=fill or (0, 0, 0, 255), font=font)
            return
        elif tag not in {"svg", "g", "a", "symbol"}:
            raise WordRenderError(f"SVG图片包含不支持的{tag}元素。")
        for child in list(element):
            walk(child, paint)

    walk(root, {"fill": "#000000"})
    output = BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue(), image.size


def _image_to_png(source: str) -> tuple[bytes, tuple[int, int]]:
    mime_type, data = _data_uri_bytes(source)
    if mime_type == "image/svg+xml":
        return _svg_to_png(data)
    expected_format = _RASTER_MIME_TYPES.get(mime_type)
    if expected_format is None:
        raise WordRenderError(f"不支持的图片格式：{mime_type or '未声明'}。")
    return _raster_to_png(data, expected_format)


class _WordHtmlRenderer:
    def __init__(self, document: DocumentType, chart_specs: Mapping[str, Any] | None) -> None:
        self.document = document
        self.chart_specs = dict(chart_specs or {})
        self.rendered_blocks = 0

    def render(self, root: Tag) -> None:
        self._render_children(root)
        if self.rendered_blocks == 0:
            raise WordRenderError("HTML正文未包含可导出的内容。")

    def _render_children(self, parent: Tag) -> None:
        pending_text: list[str] = []

        def flush_text() -> None:
            value = _clean_text("".join(pending_text))
            pending_text.clear()
            if value:
                paragraph = self.document.add_paragraph()
                run = paragraph.add_run(value)
                _set_font(run)
                self.rendered_blocks += 1

        for child in parent.children:
            if isinstance(child, NavigableString):
                pending_text.append(str(child))
                continue
            if not isinstance(child, Tag):
                continue
            flush_text()
            self._render_block(child)
        flush_text()

    def _render_block(self, node: Tag) -> None:
        tag = node.name.lower()
        if tag in _IGNORED_TAGS or node.has_attr("hidden") or str(node.get("aria-hidden") or "").lower() == "true":
            return
        style = _parse_style(node.get("style"))
        if style.get("display", "").lower() == "none" or style.get("visibility", "").lower() == "hidden":
            return
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            style_name = {"h1": "Title", "h2": "Heading 1", "h3": "Heading 2"}.get(tag, "Heading 3")
            paragraph = self.document.add_paragraph(style=style_name)
            self._format_paragraph(paragraph, node)
            self._render_inline(paragraph, node, _context_for(node, {}))
            paragraph.paragraph_format.keep_with_next = True
            self.rendered_blocks += 1
        elif tag in {"p", "address", "blockquote", "pre", "dt", "dd", "figcaption"}:
            style_name = "Caption" if tag == "figcaption" else None
            paragraph = self.document.add_paragraph(style=style_name)
            self._format_paragraph(paragraph, node)
            context = _context_for(node, {"bold": tag == "dt"})
            self._render_inline(paragraph, node, context)
            if tag in {"blockquote", "dd"}:
                paragraph.paragraph_format.left_indent = Mm(8)
            if tag == "figcaption":
                paragraph.paragraph_format.keep_with_next = True
            self.rendered_blocks += 1
        elif tag in {"ul", "ol"}:
            self._render_list(node, level=0, ordered=tag == "ol")
        elif (tag == "div" and "metric-strip" in set(node.get("class") or ())) or (
            tag == "dl" and {"card-metric-grid", "card-contract-grid"} & set(node.get("class") or ())
        ):
            self._render_metric_grid(node)
        elif tag == "table":
            self._render_table(node)
        elif tag == "figure":
            self._render_figure(node)
        elif tag == "img":
            self._render_image(node)
        elif tag == "math":
            paragraph = self.document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph._p.append(_math_element(node, display=True))
            paragraph.paragraph_format.keep_together = True
            self.rendered_blocks += 1
        elif tag == "div" and "chart" in set(node.get("class") or ()):
            self._render_chart(node)
        elif tag == "hr":
            return
        else:
            classes = set(node.get("class") or ())
            if classes & {"report-title", "document-title", "brief-title"} and _clean_text(node.get_text(" ", strip=True)):
                paragraph = self.document.add_paragraph(style="Title")
                self._format_paragraph(paragraph, node)
                self._render_inline(paragraph, node, _context_for(node, {}))
                self.rendered_blocks += 1
            else:
                self._render_children(node)

    def _format_paragraph(self, paragraph: Any, node: Tag) -> None:
        alignment = _paragraph_alignment(node)
        if alignment is not None:
            paragraph.alignment = alignment
        style = _parse_style(node.get("style"))
        before = _parse_points(style.get("margin-top"))
        after = _parse_points(style.get("margin-bottom"))
        line_height = _parse_points(style.get("line-height"))
        if before is not None:
            paragraph.paragraph_format.space_before = Pt(before)
        if after is not None:
            paragraph.paragraph_format.space_after = Pt(after)
        if line_height is not None:
            paragraph.paragraph_format.line_spacing = Pt(line_height)
        if style.get("page-break-before", "").lower() == "always" or "page-break-before" in set(node.get("class") or ()):
            paragraph.paragraph_format.page_break_before = True

    def _render_inline(self, paragraph: Any, parent: Tag, context: Mapping[str, Any]) -> None:
        for child in parent.children:
            if isinstance(child, NavigableString):
                value = _normalise_space(str(child))
                if not value:
                    continue
                if not paragraph.text and not list(paragraph._p.xpath("./m:oMath | ./m:oMathPara")):
                    value = value.lstrip()
                if value:
                    run = paragraph.add_run(value)
                    _apply_run_context(run, context)
                continue
            if not isinstance(child, Tag):
                continue
            tag = child.name.lower()
            if tag in _IGNORED_TAGS:
                continue
            if tag == "br":
                paragraph.add_run().add_break()
            elif tag == "math":
                display = str(child.get("display") or "").lower() == "block"
                paragraph._p.append(_math_element(child, display=display))
            elif tag == "img":
                self._render_image(child, paragraph=paragraph)
            elif tag in _BLOCK_TAGS:
                text = _clean_text(child.get_text(" ", strip=True))
                if text:
                    run = paragraph.add_run(text)
                    _apply_run_context(run, _context_for(child, context))
            else:
                self._render_inline(paragraph, child, _context_for(child, context))

    def _render_list(self, node: Tag, *, level: int, ordered: bool) -> None:
        items = [child for child in node.find_all("li", recursive=False)]
        for item in items:
            style_base = "List Number" if ordered else "List Bullet"
            style_name = style_base if level == 0 else f"{style_base} {min(level + 1, 3)}"
            try:
                paragraph = self.document.add_paragraph(style=style_name)
            except KeyError:
                paragraph = self.document.add_paragraph(style=style_base)
                paragraph.paragraph_format.left_indent = Mm(6 * level)
            context = _context_for(item, {})
            for child in item.children:
                if isinstance(child, Tag) and child.name.lower() in {"ul", "ol"}:
                    continue
                if isinstance(child, NavigableString):
                    value = _normalise_space(str(child))
                    if value:
                        run = paragraph.add_run(value.lstrip() if not paragraph.text else value)
                        _apply_run_context(run, context)
                elif isinstance(child, Tag):
                    self._render_inline(paragraph, child, _context_for(child, context))
            paragraph.paragraph_format.keep_together = True
            self.rendered_blocks += 1
            for nested in item.find_all(["ul", "ol"], recursive=False):
                self._render_list(nested, level=level + 1, ordered=nested.name.lower() == "ol")

    def _render_metric_grid(self, node: Tag) -> None:
        """Keep each KPI's label, value and note in one editable Word cell."""

        def visible(part: Any) -> bool:
            if not isinstance(part, Tag):
                return True
            style = _parse_style(part.get("style"))
            return not (part.name.lower() in _IGNORED_TAGS or part.has_attr("hidden")
                        or str(part.get("aria-hidden") or "").lower() == "true"
                        or style.get("display", "").lower() == "none"
                        or style.get("visibility", "").lower() == "hidden")

        is_strip = "metric-strip" in set(node.get("class") or ())
        item_classes = {"metric"} if is_strip else {"card-metric", "card-contract"}
        items: list[list[tuple[Any, bool]]] = []
        for child in node.children:
            if isinstance(child, NavigableString) and not _clean_text(str(child)):
                continue
            if not isinstance(child, Tag) or not item_classes & set(child.get("class") or ()):
                # A manually restructured block must retain all of its text.
                self._render_children(node)
                return
            if not visible(child):
                continue
            if is_strip:
                fields = [child.find(class_=name, recursive=False) for name in (
                    "metric__label", "metric__value", "metric__note")]
            else:
                fields = [child.find(name, recursive=False) for name in ("dt", "dd", "small")]
            if fields[0] is None or fields[1] is None:
                self._render_children(node)
                return
            parts = [(field, index == 0) for index, field in enumerate(fields) if field is not None]
            known = {id(field) for field, _ in parts}
            parts.extend((extra, False) for extra in child.children if id(extra) not in known
                         and (isinstance(extra, Tag) or _clean_text(str(extra))))
            parts = [(part, bold) for part, bold in parts if visible(part)]
            if parts:
                items.append(parts)
        if not items:
            return
        columns = min(2, len(items))
        row_count = (len(items) + columns - 1) // columns
        table = self.document.add_table(rows=row_count, cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Table Grid"
        section = self.document.sections[-1]
        width = section.page_width - section.left_margin - section.right_margin
        for column in table.columns:
            column.width = int(width / columns)
        for row in table.rows:
            for cell in row.cells:
                cell.width = int(width / columns)
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if len(items) % columns:
            table.cell(row_count - 1, 0).merge(table.cell(row_count - 1, columns - 1))
        _set_table_borders(table)
        for index, parts in enumerate(items):
            cell = table.cell(index // columns, index % columns)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            _set_cell_margins(cell, top=90, bottom=90, start=110, end=110)
            for part_index, (part, bold) in enumerate(parts):
                paragraph = cell.paragraphs[0] if part_index == 0 else cell.add_paragraph()
                _clear_paragraph(paragraph)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.keep_together = True
                # The row keeps its own fields together. Paragraph-level
                # keepNext would chain every row and move the whole grid.
                paragraph.paragraph_format.keep_with_next = False
                context = {"bold": bold, "size": 10.5}
                if isinstance(part, Tag):
                    self._render_inline(paragraph, part, _context_for(part, context))
                else:
                    _apply_run_context(paragraph.add_run(_clean_text(str(part))), context)
        spacer = self.document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(1)
        self.rendered_blocks += 1

    def _render_table(self, node: Tag) -> None:
        caption_node = node.find("caption", recursive=False)
        if caption_node is not None and _clean_text(caption_node.get_text(" ", strip=True)):
            caption = self.document.add_paragraph(style="Caption")
            caption.paragraph_format.keep_with_next = True
            caption_run = caption.add_run(_clean_text(caption_node.get_text(" ", strip=True)))
            caption_run.bold = True
            _set_font(caption_run, east_asia=_HEADING_FONT_EAST_ASIA)
            self.rendered_blocks += 1
        row_nodes = [row for row in node.find_all("tr") if row.find_parent("table") is node]
        if not row_nodes:
            return
        placements: list[tuple[int, int, int, int, Tag]] = []
        occupied: set[tuple[int, int]] = set()
        column_count = 0
        header_rows: set[int] = set()
        for row_index, row_node in enumerate(row_nodes):
            if row_node.find_parent("thead") is not None or all(cell.name.lower() == "th" for cell in row_node.find_all(["th", "td"], recursive=False)):
                header_rows.add(row_index)
            column_index = 0
            for cell_node in row_node.find_all(["th", "td"], recursive=False):
                while (row_index, column_index) in occupied:
                    column_index += 1
                row_span = _positive_int(cell_node.get("rowspan"))
                column_span = _positive_int(cell_node.get("colspan"))
                placements.append((row_index, column_index, row_span, column_span, cell_node))
                for target_row in range(row_index, row_index + row_span):
                    for target_column in range(column_index, column_index + column_span):
                        occupied.add((target_row, target_column))
                column_index += column_span
                column_count = max(column_count, column_index)
        if column_count == 0:
            return
        table = self.document.add_table(rows=len(row_nodes), cols=column_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        table.style = "Table Grid"
        _set_table_borders(table)
        for row_index, column_index, row_span, column_span, _cell_node in placements:
            if row_span > 1 or column_span > 1:
                table.cell(row_index, column_index).merge(table.cell(row_index + row_span - 1, column_index + column_span - 1))
        for row_index, column_index, _row_span, _column_span, cell_node in placements:
            cell = table.cell(row_index, column_index)
            self._render_cell(cell, cell_node)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            style = _parse_style(cell_node.get("style"))
            background = style.get("background-color") or style.get("background")
            parsed_background = _parse_color(background)
            if parsed_background is not None:
                _set_cell_shading(cell, str(parsed_background))
            elif row_index in header_rows:
                _set_cell_shading(cell, "315D8A")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                _set_cell_shading(cell, "F4F7FA")
        for row_index in sorted(header_rows):
            _set_repeat_table_header(table.rows[row_index])
        # Keep each record together, without chaining separate rows. Otherwise
        # Word may put a cell's trailing padding alone on the next page and
        # repeat the header above an apparently empty continuation row.
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        spacer = self.document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(1)
        self.rendered_blocks += 1

    def _render_cell(self, cell: Any, node: Tag) -> None:
        paragraph = cell.paragraphs[0]
        _clear_paragraph(paragraph)
        alignment = _paragraph_alignment(node)
        if alignment is not None:
            paragraph.alignment = alignment
        block_children = [child for child in node.children if isinstance(child, Tag) and child.name.lower() in _BLOCK_TAGS]
        if not block_children:
            self._render_inline(paragraph, node, _context_for(node, {"bold": node.name.lower() == "th", "size": 9.5}))
            return
        first = True
        for child in node.children:
            if isinstance(child, NavigableString):
                text = _clean_text(str(child))
                if text:
                    target = paragraph if first else cell.add_paragraph()
                    run = target.add_run(text)
                    _apply_run_context(run, {"bold": node.name.lower() == "th", "size": 9.5})
                    first = False
                continue
            if not isinstance(child, Tag):
                continue
            if child.name.lower() in {"p", "div", "figcaption"}:
                target = paragraph if first else cell.add_paragraph()
                self._format_paragraph(target, child)
                self._render_inline(target, child, _context_for(child, {"bold": node.name.lower() == "th", "size": 9.5}))
                first = False
            elif child.name.lower() in {"ul", "ol"}:
                for item in child.find_all("li", recursive=False):
                    target = paragraph if first else cell.add_paragraph()
                    target.style = "List Number" if child.name.lower() == "ol" else "List Bullet"
                    self._render_inline(target, item, _context_for(item, {"size": 9.5}))
                    first = False
            elif child.name.lower() == "math":
                target = paragraph if first else cell.add_paragraph()
                target._p.append(_math_element(child, display=str(child.get("display") or "").lower() == "block"))
                first = False
            else:
                target = paragraph if first else cell.add_paragraph()
                self._render_inline(target, child, _context_for(child, {"bold": node.name.lower() == "th", "size": 9.5}))
                first = False

    def _render_figure(self, node: Tag) -> None:
        chart = next((child for child in node.find_all("div") if "chart" in set(child.get("class") or ())), None)
        if chart is not None:
            spec = self._chart_spec(chart)
            caption = node.find("figcaption")
            title = _clean_text(caption.get_text(" ", strip=True)) if caption is not None else str(spec.get("title") or "图表")
            data_container = next((child for child in node.find_all("div") if "chart-data" in set(child.get("class") or ())), None)
            add_data_table = data_container is None or not self._chart_data_matches(data_container, spec)
            self._add_chart(spec, title=title, add_data_table=add_data_table)
            for child in node.children:
                if not isinstance(child, Tag) or child is caption or child is chart:
                    continue
                self._render_block(child)
            return
        images = node.find_all("img", recursive=False)
        for image in images:
            self._render_image(image)
        for child in node.children:
            if isinstance(child, Tag) and child.name.lower() != "img":
                self._render_block(child)

    def _render_image(self, node: Tag, paragraph: Any | None = None) -> None:
        source = str(node.get("src") or "")
        png, size = _image_to_png(source)
        target = paragraph or self.document.add_paragraph()
        if paragraph is None:
            target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        width_inches = min(_CONTENT_WIDTH_INCHES, max(0.35, size[0] / 144.0))
        height_inches = size[1] / max(1, size[0]) * width_inches
        if height_inches > 7.7:
            width_inches *= 7.7 / height_inches
        run = target.add_run()
        shape = run.add_picture(BytesIO(png), width=Inches(width_inches))
        alt = _clean_text(str(node.get("alt") or ""))
        if alt:
            shape._inline.docPr.set("descr", alt)
            shape._inline.docPr.set("title", alt)
        target.paragraph_format.keep_together = True
        self.rendered_blocks += 1

    def _chart_spec(self, node: Tag) -> Mapping[str, Any]:
        chart_id = _clean_text(str(node.get("id") or ""))
        if not chart_id:
            raise WordRenderError("Word图表缺少稳定id。")
        spec = self.chart_specs.get(chart_id)
        if not isinstance(spec, Mapping):
            raise WordRenderError(f"报告图表{chart_id}缺少当前Designer图表规格。")
        return spec

    def _render_chart(self, node: Tag) -> None:
        spec = self._chart_spec(node)
        self._add_chart(spec, title=str(spec.get("title") or "图表"), add_data_table=True)

    def _add_chart(self, spec: Mapping[str, Any], *, title: str, add_data_table: bool) -> None:
        caption = self.document.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.keep_with_next = True
        caption_run = caption.add_run(title or str(spec.get("title") or "图表"))
        caption_run.bold = True
        _set_font(caption_run, east_asia=_HEADING_FONT_EAST_ASIA)
        figure = self.document.add_paragraph()
        figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure.paragraph_format.keep_together = True
        image = figure.add_run().add_picture(BytesIO(_render_chart_png(spec)), width=Inches(_CONTENT_WIDTH_INCHES))
        image._inline.docPr.set("descr", str(spec.get("title") or title or "图表"))
        image._inline.docPr.set("title", str(spec.get("title") or title or "图表"))
        self.rendered_blocks += 2
        if add_data_table:
            self._render_chart_data_table(spec)

    def _chart_data_matches(self, container: Tag, spec: Mapping[str, Any]) -> bool:
        visible = _clean_text(container.get_text(" ", strip=True))
        expected: list[str] = [str(value) for value in list(spec.get("x") or [])]
        if str(spec.get("type") or "line").lower() == "heatmap":
            expected.extend(str(value) for value in list(spec.get("y") or []))
            expected.extend(_display_value(row[2], spec.get("value_format"), spec.get("value_suffix")) for row in list(spec.get("data") or []) if isinstance(row, (list, tuple)) and len(row) == 3)
        else:
            for series in list(spec.get("series") or []):
                if isinstance(series, Mapping):
                    expected.extend(_display_value(value, spec.get("value_format"), spec.get("value_suffix")) for value in list(series.get("data") or []))
        return all(value in visible for value in expected if value)

    def _render_chart_data_table(self, spec: Mapping[str, Any]) -> None:
        chart_type = str(spec.get("type") or "line").lower()
        x_values = list(spec.get("x") or [])
        value_format = spec.get("value_format")
        suffix = spec.get("value_suffix")
        if chart_type == "heatmap":
            y_values = list(spec.get("y") or [])
            values = {
                (int(row[0]), int(row[1])): _display_value(row[2], value_format, suffix)
                for row in list(spec.get("data") or [])
                if isinstance(row, (list, tuple)) and len(row) == 3 and str(row[0]).lstrip("-").isdigit() and str(row[1]).lstrip("-").isdigit()
            }
            rows = [[str(spec.get("y_axis_name") or "纵轴") + " / " + str(spec.get("x_axis_name") or "横轴"), *(str(value) for value in x_values)]]
            rows.extend([str(y_value), *(values.get((x_index, y_index), "—") for x_index in range(len(x_values)))] for y_index, y_value in enumerate(y_values))
        else:
            series = [item for item in list(spec.get("series") or []) if isinstance(item, Mapping)]
            rows = [[str(spec.get("x_axis_name") or "横轴"), *(str(item.get("name") or f"序列{index + 1}") for index, item in enumerate(series))]]
            for index, x_value in enumerate(x_values):
                rows.append([
                    str(x_value),
                    *(
                        _display_value(list(item.get("data") or [])[index], value_format, suffix)
                        if index < len(list(item.get("data") or [])) else "—"
                        for item in series
                    ),
                ])
        self._add_matrix_table(rows)

    def _add_matrix_table(self, rows: Sequence[Sequence[str]]) -> None:
        if not rows or not rows[0]:
            return
        table = self.document.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        _set_table_borders(table)
        for row_index, values in enumerate(rows):
            for column_index, value in enumerate(values):
                cell = table.cell(row_index, column_index)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _set_cell_margins(cell)
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index else WD_ALIGN_PARAGRAPH.LEFT
                run = paragraph.add_run(str(value))
                _set_font(run)
                run.font.size = Pt(9)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    _set_cell_shading(cell, "315D8A")
                elif row_index % 2 == 0:
                    _set_cell_shading(cell, "F4F7FA")
        _set_repeat_table_header(table.rows[0])
        spacer = self.document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(1)
        self.rendered_blocks += 1


def _select_document_root(soup: BeautifulSoup) -> Tag:
    preferred = soup.select_one("#report-main.report-document")
    if isinstance(preferred, Tag):
        return preferred
    main = soup.find("main")
    if isinstance(main, Tag):
        return main
    if isinstance(soup.body, Tag):
        return soup.body
    return soup


def render_docx(html_content: str, *, chart_specs: Mapping | None = None) -> bytes:
    """Convert cleaned Designer HTML to an editable A4 DOCX byte stream.

    ``#report-main.report-document`` is authoritative when present; otherwise
    the first ``main`` element and then ``body`` are used.  ``chart_specs`` is
    the current Designer mapping from chart id to frozen line, bar or heatmap
    specification.
    """

    if not isinstance(html_content, str) or not html_content.strip():
        raise WordRenderError("Word输入HTML不能为空。")
    if chart_specs is not None and not isinstance(chart_specs, Mapping):
        raise WordRenderError("chart_specs必须是图表id到Designer规格的映射。")
    soup = BeautifulSoup(html_content, "html.parser")
    root = _select_document_root(soup)
    document = Document()
    _configure_document(document)
    heading = soup.find("title") or root.find("h1")
    document.core_properties.title = _clean_text(heading.get_text(" ", strip=True)) if heading else "研究报告"
    _WordHtmlRenderer(document, chart_specs).render(root)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


__all__ = ["WordRenderError", "render_docx"]
